"""Tests for artifact_synthesis — DOCX/XLSX/PPTX/CSV/PDF render paths.

These guard the user-visible contract: when the Document Generator agent
emits an `[[ARTIFACT name="x.docx"]] ... [[/ARTIFACT]]` block, the bytes
returned by `synthesize()` must open in the target application without
errors.
"""
from __future__ import annotations

from io import BytesIO

import pytest


_SAMPLE_MD = """# Quarterly Review

A short narrative intro paragraph that should land on the Notes sheet
when this becomes a spreadsheet.

## Budget table

| Category | Budget | Actual | Variance |
|----------|-------:|-------:|---------:|
| Marketing | 50000 | 48200 | -1800 |
| Engineering | 200000 | 211500 | 11500 |
| Sales | 100000 | 95000 | -5000 |

## Headcount table

| Team | Q2 | Q3 |
|------|---:|---:|
| Engineering | 6 | 8 |
| Sales | 4 | 5 |

Closing thoughts in another paragraph.
"""


def test_docx_renders_with_real_paragraphs_and_tables() -> None:
    """User asks for a Word doc → must open as a real .docx with the
    headings and tables intact."""
    from artifact_synthesis import synthesize
    data, mime = synthesize(_SAMPLE_MD, "docx", title="Quarterly Review")
    assert mime.endswith("wordprocessingml.document")
    assert len(data) > 1000
    from docx import Document
    doc = Document(BytesIO(data))
    assert len(doc.paragraphs) > 0
    # 2 markdown tables → 2 docx tables
    assert len(doc.tables) == 2


def test_xlsx_renders_one_sheet_per_table_plus_notes() -> None:
    """Multi-table Markdown → multi-sheet XLSX with prose preserved
    on a Notes sheet."""
    from artifact_synthesis import synthesize
    data, mime = synthesize(_SAMPLE_MD, "xlsx", title="Quarterly Review")
    assert mime.endswith("spreadsheetml.sheet")
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(data))
    # Notes always first, then Sheet1..N for each table.
    assert wb.sheetnames[0] == "Notes"
    assert "Sheet1" in wb.sheetnames
    assert "Sheet2" in wb.sheetnames
    # Header row should be present in Sheet1.
    assert wb["Sheet1"].cell(1, 1).value == "Category"
    # Freeze pane on data sheets.
    assert wb["Sheet1"].freeze_panes == "A2"


def test_pptx_each_heading_becomes_a_slide() -> None:
    """Markdown `#` headings drive slide structure."""
    from artifact_synthesis import synthesize
    data, mime = synthesize(_SAMPLE_MD, "pptx", title="Quarterly Review")
    assert mime.endswith("presentationml.presentation")
    from pptx import Presentation
    prs = Presentation(BytesIO(data))
    # Title slide + 3 content slides (Quarterly Review, Budget table, Headcount table)
    assert len(prs.slides) >= 3
    # First slide is the title; subsequent slides have non-empty titles.
    titles = [s.shapes.title.text for s in prs.slides if s.shapes.title]
    assert any("Quarterly Review" in t for t in titles)
    assert any("Budget table" in t for t in titles)


def test_csv_picks_first_table_and_writes_bom() -> None:
    """CSV path picks the first markdown table; bytes are BOM-prefixed
    so Excel auto-detects UTF-8."""
    from artifact_synthesis import synthesize
    data, mime = synthesize(_SAMPLE_MD, "csv", title="Quarterly Review")
    assert mime == "text/csv"
    assert data.startswith(b"\xef\xbb\xbf"), "Missing UTF-8 BOM"
    text = data.decode("utf-8-sig")
    first_line = text.splitlines()[0]
    assert "Category" in first_line and "Budget" in first_line


def test_pdf_renders_via_reportlab_fallback() -> None:
    """PDF path always returns valid bytes (reportlab fallback if the
    branded renderer isn't available)."""
    from artifact_synthesis import synthesize
    data, mime = synthesize(_SAMPLE_MD, "pdf", title="Quarterly Review")
    assert mime == "application/pdf"
    assert data[:4] == b"%PDF"


def test_unsupported_extension_raises() -> None:
    from artifact_synthesis import synthesize
    with pytest.raises(ValueError):
        synthesize("body", "rtf")


def test_supported_set_contains_all_five() -> None:
    from artifact_synthesis import SUPPORTED_BINARY_EXTS
    assert SUPPORTED_BINARY_EXTS == {"pdf", "docx", "xlsx", "pptx", "csv"}


def test_xlsx_with_no_tables_still_produces_notes_sheet() -> None:
    """Edge case: agent emits a narrative-only body with no tables.
    The XLSX path should still produce a valid workbook with the prose
    on the Notes sheet — not crash, not produce an empty file."""
    from artifact_synthesis import synthesize
    data, mime = synthesize(
        "Just a single paragraph of prose with no tables in sight.",
        "xlsx",
    )
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(data))
    assert "Notes" in wb.sheetnames


def test_pptx_with_no_headings_still_produces_at_least_two_slides() -> None:
    """Edge case: agent emits a body with no `#` headings. The PPTX
    path should produce a title slide + a Content slide with the body
    dumped in — not a 1-slide empty deck."""
    from artifact_synthesis import synthesize
    data, mime = synthesize(
        "Just narrative prose with bullets:\n- one\n- two\n- three",
        "pptx",
        title="My Deck",
    )
    from pptx import Presentation
    prs = Presentation(BytesIO(data))
    assert len(prs.slides) >= 2
